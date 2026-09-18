"""
Script para generar los entregables oficiales DAS.docx y DAS.pdf
a partir del contenido del Documento de Arquitectura de Software.
"""
import os
import re
import html
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls

from reportlab.lib.pagesizes import letter
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, KeepTogether, HRFlowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors

DOCS_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'docs'))
DAS_MD = os.path.join(DOCS_DIR, 'DAS.md')
DAS_DOCX = os.path.join(DOCS_DIR, 'DAS.docx')
DAS_PDF = os.path.join(DOCS_DIR, 'DAS.pdf')

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def md_to_reportlab(text):
    if not text:
        return ""
    # Normalizar saltos de línea
    text = text.replace('<br>', '___BR___')
    # Escapar caracteres HTML básicos
    text = html.escape(text, quote=False)
    text = text.replace('___BR___', '<br/>')

    # Convertir **negrita** primero
    text = re.sub(r'\*\*(.+?)\*\*', r'<b>\1</b>', text)

    # Convertir `código` (limpiar backticks sin tags conflictivos)
    text = re.sub(r'`(.+?)`', r'<b>\1</b>', text)

    # Cursiva sólo cuando sea claramente una palabra o frase, evitando asteriscos como wildcard (*.md)
    text = re.sub(r'(?<![\w\*\-])\*([a-zA-ZáéíóúÁÉÍÓÚñÑ ]+?)\*(?![\w\*\-])', r'<i>\1</i>', text)

    # Limpiar posibles caracteres matemáticos o símbolos problemáticos
    text = text.replace(r'\le', '&le;').replace(r'\ge', '&ge;').replace(r'\rightarrow', '&rarr;')
    text = text.replace('$', '')
    return text

def generar_docx():
    print("[DOCX] Generando DAS.docx...")
    doc = Document()

    # Configurar márgenes
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # Portada / Encabezado
    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("Documento de Arquitectura de Software (DAS / SAD)\n")
    run_title.bold = True
    run_title.font.size = Pt(22)
    run_title.font.color.rgb = RGBColor(0, 51, 102)

    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Plantilla oficial para proyectos individuales — arc42 simplificado + Modelo C4 + ADRs + Corte Vertical\n")
    run_sub.italic = True
    run_sub.font.size = Pt(12)
    run_sub.font.color.rgb = RGBColor(100, 116, 139)

    # Metadatos del proyecto en tabla
    meta_table = doc.add_table(rows=5, cols=2)
    meta_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Curso:", "Arquitectura de Software / Ingeniería de Software"),
        ("Modalidad:", "Proyecto individual"),
        ("Estudiante:", "Juan Pérez (Estudiante de Ingeniería de Software)"),
        ("Proyecto:", "Sistema de Gestión y Emisión Digital de Trámites y Certificados 24/7 (Caso 2: Alcaldía Municipal)"),
        ("Versión y Fecha:", "Versión 1.0 — 17 de Septiembre de 2026")
    ]
    for i, (k, v) in enumerate(meta_data):
        row = meta_table.rows[i]
        c0, c1 = row.cells[0], row.cells[1]
        c0.paragraphs[0].add_run(k).bold = True
        c1.paragraphs[0].add_run(v)
        set_cell_background(c0, "F1F5F9")
        set_cell_background(c1, "FFFFFF")

    doc.add_paragraph("\n")

    with open(DAS_MD, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    in_code_block = False
    in_table = False
    table_lines = []

    def flush_table(lines_to_parse):
        if not lines_to_parse:
            return
        rows_data = []
        for l in lines_to_parse:
            raw = l.strip()
            if not raw.startswith('|'):
                continue
            cols = [c.strip() for c in raw.split('|')[1:-1]]
            if all(set(c).issubset({'-', ':', ' '}) for c in cols):
                continue
            rows_data.append(cols)

        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        t = doc.add_table(rows=len(rows_data), cols=cols_count)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        for r_idx, row in enumerate(rows_data):
            for c_idx, val in enumerate(row):
                if c_idx < cols_count:
                    cell = t.cell(r_idx, c_idx)
                    clean_val = val.replace('<br>', '\n').replace('**', '').replace('`', '')
                    p = cell.paragraphs[0]
                    r = p.add_run(clean_val)
                    if r_idx == 0:
                        r.bold = True
                        set_cell_background(cell, "003366")
                        r.font.color.rgb = RGBColor(255, 255, 255)
                    else:
                        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
                        set_cell_background(cell, bg)
        doc.add_paragraph()

    skip_header = True
    for line in lines:
        stripped = line.strip()

        if "## SECCIÓN 1" in line:
            skip_header = False

        if skip_header:
            continue

        if stripped.startswith('```'):
            in_code_block = not in_code_block
            continue

        if in_code_block:
            p = doc.add_paragraph()
            r = p.add_run(stripped)
            r.font.name = 'Consolas'
            r.font.size = Pt(9)
            p.paragraph_format.left_indent = Inches(0.4)
            continue

        if stripped.startswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            in_table = False
            flush_table(table_lines)
            table_lines = []

        if not stripped:
            continue

        if stripped.startswith('## '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[3:])
            r.bold = True
            r.font.size = Pt(15)
            r.font.color.rgb = RGBColor(0, 51, 102)
        elif stripped.startswith('### '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[4:])
            r.bold = True
            r.font.size = Pt(13)
            r.font.color.rgb = RGBColor(0, 137, 123)
        elif stripped.startswith('#### '):
            p = doc.add_paragraph()
            r = p.add_run(stripped[5:])
            r.bold = True
            r.font.size = Pt(11)
            r.font.color.rgb = RGBColor(51, 65, 85)
        elif stripped.startswith('> '):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.4)
            r = p.add_run(stripped[2:].replace('**', '').replace('`', ''))
            r.italic = True
            r.font.color.rgb = RGBColor(30, 41, 59)
        elif stripped.startswith('- ') or stripped.startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            text = stripped[2:].replace('**', '').replace('`', '')
            p.add_run(text)
        elif re.match(r'^\d+\.\s', stripped):
            p = doc.add_paragraph(style='List Number')
            text = re.sub(r'^\d+\.\s', '', stripped).replace('**', '').replace('`', '')
            p.add_run(text)
        else:
            p = doc.add_paragraph()
            text = stripped.replace('**', '').replace('`', '')
            p.add_run(text)

    if in_table and table_lines:
        flush_table(table_lines)

    doc.save(DAS_DOCX)
    print(f"[DOCX] Archivo generado exitosamente en: {DAS_DOCX}")

def generar_pdf():
    print("[PDF] Generando DAS.pdf...")
    pdf = SimpleDocTemplate(
        DAS_PDF,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        leading=22,
        textColor=colors.HexColor('#003366'),
        alignment=1,
        spaceAfter=6
    )
    subtitle_style = ParagraphStyle(
        'DocSubtitle',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=10,
        leading=13,
        textColor=colors.HexColor('#64748B'),
        alignment=1,
        spaceAfter=15
    )
    h1_style = ParagraphStyle(
        'DocH1',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=13,
        leading=16,
        textColor=colors.HexColor('#003366'),
        spaceBefore=14,
        spaceAfter=6
    )
    h2_style = ParagraphStyle(
        'DocH2',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=14,
        textColor=colors.HexColor('#00897B'),
        spaceBefore=10,
        spaceAfter=4
    )
    h3_style = ParagraphStyle(
        'DocH3',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=9.5,
        leading=12,
        textColor=colors.HexColor('#1E293B'),
        spaceBefore=8,
        spaceAfter=3
    )
    body_style = ParagraphStyle(
        'DocBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#1E293B'),
        spaceAfter=5
    )
    quote_style = ParagraphStyle(
        'DocQuote',
        parent=styles['Normal'],
        fontName='Helvetica-Oblique',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#0F172A'),
        leftIndent=15,
        spaceAfter=5
    )
    table_header_style = ParagraphStyle(
        'DocTH',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=10,
        textColor=colors.white,
        alignment=0
    )
    table_cell_style = ParagraphStyle(
        'DocTD',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#1E293B')
    )

    story = []

    # Título y Subtítulo
    story.append(Paragraph("Documento de Arquitectura de Software (DAS / SAD)", title_style))
    story.append(Paragraph("Plantilla oficial para proyectos individuales — arc42 simplificado + Modelo C4 + ADRs + Corte Vertical", subtitle_style))

    # Metadatos del proyecto
    meta_table_data = [
        [Paragraph("<b>Curso:</b>", body_style), Paragraph("Arquitectura de Software / Ingeniería de Software", body_style)],
        [Paragraph("<b>Modalidad:</b>", body_style), Paragraph("Proyecto individual", body_style)],
        [Paragraph("<b>Estudiante:</b>", body_style), Paragraph("Juan Pérez (Estudiante de Ingeniería de Software)", body_style)],
        [Paragraph("<b>Proyecto:</b>", body_style), Paragraph("Sistema de Gestión y Emisión Digital de Trámites y Certificados 24/7 (Caso 2)", body_style)],
        [Paragraph("<b>Versión / Fecha:</b>", body_style), Paragraph("Versión 1.0 — 17 de Septiembre de 2026", body_style)]
    ]
    t_meta = Table(meta_table_data, colWidths=[110, 420])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F1F5F9')),
        ('BACKGROUND', (1,0), (1,-1), colors.white),
        ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
        ('TOPPADDING', (0,0), (-1,-1), 3),
        ('BOTTOMPADDING', (0,0), (-1,-1), 3),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 10))

    with open(DAS_MD, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    skip_header = True
    in_code = False
    in_table = False
    table_lines = []

    def flush_pdf_table(lines_to_parse):
        if not lines_to_parse:
            return
        rows_data = []
        for l in lines_to_parse:
            raw = l.strip()
            if not raw.startswith('|'):
                continue
            cols = [c.strip() for c in raw.split('|')[1:-1]]
            if all(set(c).issubset({'-', ':', ' '}) for c in cols):
                continue
            rows_data.append(cols)

        if not rows_data:
            return

        cols_count = max(len(r) for r in rows_data)
        available_width = 530
        col_width = available_width / cols_count

        table_formatted = []
        for r_idx, r in enumerate(rows_data):
            row_cells = []
            for val in r:
                formatted = md_to_reportlab(val)
                st = table_header_style if r_idx == 0 else table_cell_style
                row_cells.append(Paragraph(formatted, st))
            while len(row_cells) < cols_count:
                row_cells.append(Paragraph("", table_cell_style))
            table_formatted.append(row_cells)

        t_elem = Table(table_formatted, colWidths=[col_width]*cols_count)
        t_elem.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#003366')),
            ('ALIGN', (0,0), (-1,-1), 'LEFT'),
            ('VALIGN', (0,0), (-1,-1), 'TOP'),
            ('BOX', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('INNERGRID', (0,0), (-1,-1), 0.5, colors.HexColor('#E2E8F0')),
            ('ROWBACKGROUNDS', (0,1), (-1,-1), [colors.white, colors.HexColor('#F8FAFC')]),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(t_elem)
        story.append(Spacer(1, 8))

    for line in lines:
        stripped = line.strip()

        if "## SECCIÓN 1" in line:
            skip_header = False

        if skip_header:
            continue

        if stripped.startswith('```'):
            in_code = not in_code
            continue

        if in_code:
            continue

        if stripped.startswith('|'):
            in_table = True
            table_lines.append(line)
            continue
        elif in_table:
            in_table = False
            flush_pdf_table(table_lines)
            table_lines = []

        if not stripped:
            continue

        if stripped.startswith('## '):
            story.append(Paragraph(stripped[3:], h1_style))
            story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#CBD5E1'), spaceAfter=5))
        elif stripped.startswith('### '):
            story.append(Paragraph(stripped[4:], h2_style))
        elif stripped.startswith('#### '):
            story.append(Paragraph(stripped[5:], h3_style))
        elif stripped.startswith('> '):
            story.append(Paragraph(md_to_reportlab(stripped[2:]), quote_style))
        elif stripped.startswith('- ') or stripped.startswith('* '):
            formatted = md_to_reportlab(stripped[2:])
            story.append(Paragraph(f"• {formatted}", body_style))
        elif re.match(r'^\d+\.\s', stripped):
            formatted = md_to_reportlab(stripped)
            story.append(Paragraph(formatted, body_style))
        else:
            formatted = md_to_reportlab(stripped)
            story.append(Paragraph(formatted, body_style))

    if in_table and table_lines:
        flush_pdf_table(table_lines)

    pdf.build(story)
    print(f"[PDF] Archivo generado exitosamente en: {DAS_PDF}")

if __name__ == '__main__':
    generar_docx()
    generar_pdf()
